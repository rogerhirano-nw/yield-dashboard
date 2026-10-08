"""Build the site-served campaign GA4 UTM tracking-URL template (deliverables/site_served_campaign_utm_urls.xlsx).

Run from deliverables/: python ../scripts/build_utm_sheet.py
"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter as L
F="Arial"; wb=Workbook(); ws=wb.active; ws.title="Tracking URLs"
yellow=PatternFill("solid",start_color="FFF2CC"); head=PatternFill("solid",start_color="1F1E19")
thin=Side(style="thin",color="BFBFBF"); bd=Border(top=thin,bottom=thin,left=thin,right=thin)
blue=Font(name=F,color="0000FF")
ws["A1"]="Site-served campaign — GA4 tracking URLs"; ws["A1"].font=Font(name=F,size=14,bold=True)
ws["A2"]=("Built per Google's GA4 URL-builder guidance (support.google.com/analytics/answer/10917952). "
          "Yellow cells are inputs; the last column builds the click-through URL. Blank row inputs fall back to the campaign defaults.")
ws["A2"].font=Font(name=F,size=9,italic=True)
defaults=[("Advertiser","Example Brand",""),
 ("Default landing page URL","https://www.example.com/offer","Advertiser's destination; a URL that already has a ? gets & appended."),
 ("utm_id","","Campaign ID — optional; use the same ID as any GA4 campaign-data upload (e.g. the IO/SO number)."),
 ("utm_source","newsweek","Referrer."),
 ("utm_medium","display","Marketing medium (banner/display)."),
 ("utm_campaign","example_brand_fall_2026","Product, slogan or promo."),
 ("utm_source_platform","Google Ad Manager","Platform directing the traffic — site-served via GAM."),
 ("utm_creative_format","display","Default creative type; override per row. Not reported in GA4 yet."),
 ("utm_marketing_tactic","prospecting","Default targeting tactic; override per row. Not reported in GA4 yet.")]
ws["A4"]="Campaign defaults"; ws["A4"].font=Font(name=F,bold=True)
D={}
for i,(k,v,note) in enumerate(defaults,start=5):
    ws.cell(i,1,k).font=Font(name=F); c=ws.cell(i,2,v); c.font=blue; c.fill=yellow; c.border=bd
    ws.cell(i,3,note).font=Font(name=F,size=9,italic=True,color="595959"); D[k]=f"$B${i}"
hr=16
# (header, kind) kind: in = input, calc
cols=["#","GAM line item / placement","Ad size","Creative name","Landing page URL (blank = default)",
 "utm_creative_format (blank = default)","utm_marketing_tactic (blank = default)","utm_term (optional)",
 "utm_id","utm_source","utm_medium","utm_campaign","utm_source_platform","utm_content",
 "utm_creative_format","utm_marketing_tactic","Tracking URL (click-through)"]
inputs={2,3,4,5,6,7,8}
for j,h in enumerate(cols,1):
    c=ws.cell(hr,j,h); c.font=Font(name=F,bold=True,color="FFFFFF"); c.fill=head; c.border=bd
    c.alignment=Alignment(wrap_text=True,vertical="center")
ws.row_dimensions[hr].height=32
def n(x): return f'LOWER(SUBSTITUTE(TRIM({x})," ","_"))'
dv=DataValidation(type="list",formula1='"display,native,video,rich_media,interstitial"',allow_blank=True)
ws.add_data_validation(dv)
first,last=hr+1,hr+40
dv.add(f"F{first}:F{last}")
for r in range(first,last+1):
    g=lambda s:f'=IF($B{r}="","",{s})'
    ws[f"A{r}"]=f'=IF(B{r}="","",ROW()-{hr})'
    ws[f"I{r}"]=f'=IF(OR($B{r}="",{D["utm_id"]}=""),"",{n(D["utm_id"])})'
    ws[f"J{r}"]=g(n(D["utm_source"])); ws[f"K{r}"]=g(n(D["utm_medium"]))
    ws[f"L{r}"]=g(n(D["utm_campaign"])); ws[f"M{r}"]=g(n(D["utm_source_platform"]))
    ws[f"N{r}"]=g(n(f'IF(D{r}="",C{r},D{r})'))
    ws[f"O{r}"]=g(n(f'IF(F{r}="",{D["utm_creative_format"]},F{r})'))
    ws[f"P{r}"]=g(n(f'IF(G{r}="",{D["utm_marketing_tactic"]},G{r})'))
    lp=f'TRIM(IF(E{r}="",{D["Default landing page URL"]},E{r}))'
    def opt(name,ref): return f'&IF({ref}="","","&{name}="&{ref})'
    ws[f"Q{r}"]=(f'=IF(B{r}="","",{lp}&IF(ISNUMBER(FIND("?",{lp})),"&","?")'
        f'&"utm_source="&J{r}&"&utm_medium="&K{r}&"&utm_campaign="&L{r}'
        +opt("utm_id",f"I{r}")+opt("utm_source_platform",f"M{r}")
        +f'&IF(H{r}="","","&utm_term="&{n(f"H{r}")})'
        +opt("utm_content",f"N{r}")+opt("utm_creative_format",f"O{r}")+opt("utm_marketing_tactic",f"P{r}")+')')
    for j in range(1,len(cols)+1):
        c=ws.cell(r,j); c.border=bd
        if j in inputs: c.fill=yellow; c.font=blue
        else: c.font=Font(name=F)
for j,v in enumerate(["Homepage takeover","970x250","ExampleBrand_Fall_970x250_v1"],2): ws.cell(first,j,v)
widths=[5,30,10,32,34,16,16,14,12,12,12,24,20,30,14,14,120]
for j,w in enumerate(widths,1): ws.column_dimensions[L(j)].width=w
ws.column_dimensions["A"].width=26; ws.column_dimensions["B"].width=32; ws.column_dimensions["C"].width=12
ws.freeze_panes=f"C{first}"
ws.cell(last+2,1,"Row 17 is an example — overwrite it. Values are lowercased with spaces→underscores (GA4 is case-sensitive). "
        "utm_content = creative name, else ad size. utm_creative_format / utm_marketing_tactic are not reported in GA4 yet but are passed through.").font=Font(name=F,size=9,italic=True)
wb.save("site_served_campaign_utm_urls.xlsx")
