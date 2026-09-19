"""
Forward-looking avails for a FUTURE month, by country, Display + Video.

This is the forecast companion to `gam_intl_avails.py` (which reports
historical actuals). Use this one when the month being sold has not happened
yet — it asks GAM's own ForecastService what a prospective line item could
book, so it carries GAM's seasonality model and nets out inventory already
reserved by other line items.

Scope matches gam_intl_avails.py: Display = top-level ad unit `newsweek`
limited to 300x250 / 320x50 / 970x250 / 728x90, Video = `vid.newsweek`.
The US is excluded.

Two numbers come back per cut, and they mean different things:
  - matched   = all inventory matching the targeting (the gross pool)
  - available = what is still unreserved, i.e. what you can actually sell
Both are reported; `available` is the one to quote to a buyer.

GAM-side notes (all found empirically against network 22541732127)
------------------------------------------------------------------
- Forecasting is SOAP-only. The REST v1 surface (`google-ads-admanager`) has
  no forecast service at all, so this goes through `googleads` /
  ForecastService v202605 while the historical pull stays on REST.
- A VIDEO_PLAYER line item needs THREE things set together or the call fails:
  `environmentType=VIDEO_PLAYER`, `targeting.requestPlatformTargeting`
  (else NotNullError.NULL @ targeting.requestPlatformTargeting), and
  `videoMaxDuration` (else INVALID_MAX_VIDEO_CREATIVE_DURATION, trigger '0').
  The field is `videoMaxDuration` on the line item — NOT
  `maxVideoCreativeDuration`, which the fault message names but the v202605
  WSDL does not define, so zeep rejects it as an unknown key.
- A line item carrying several creative placeholders forecasts the UNION of
  those sizes, which is exactly the de-duplicated "eligible for at least one
  of the four" figure. Per-size calls still overlap each other and must not
  be summed.
- Countries are joined to Geo_Target by ISO country CODE, not name: the
  reporting API and the Geo_Target table disagree on wording for several
  ("The Netherlands" vs "Netherlands", "Türkiye", "Czechia").

Usage:
    python scripts/gam_avails_forecast.py                     # next month
    python scripts/gam_avails_forecast.py --month 2026-10
    python scripts/gam_avails_forecast.py --month 2026-10 --no-per-size
    python scripts/gam_avails_forecast.py --xlsx out/oct.xlsx --csv-dir out/
"""

from __future__ import annotations

import argparse
import calendar
import os
import sys
import threading
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import date, timedelta
from pathlib import Path

import pandas as pd

REPO_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO_ROOT))

env_file = REPO_ROOT / ".env"
if env_file.exists():
    for line in env_file.read_text().splitlines():
        line = line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        k, v = line.split("=", 1)
        os.environ.setdefault(k.strip(), v.strip().strip('"').strip("'"))

from gam_client import GAMClient  # noqa: E402

API_VERSION = "v202605"
DISPLAY_UNIT = "newsweek"
VIDEO_UNIT = "vid.newsweek"
DISPLAY_SIZES = [(300, 250), (320, 50), (970, 250), (728, 90)]
VIDEO_SIZE = (640, 360)
VIDEO_MAX_DURATION_MS = 30_000
EXCLUDE_CODES = {"US"}
ANY_SIZE = "ANY of the 4 (de-duplicated)"
TZ = "America/New_York"

_throttle = threading.Semaphore(3)


def _month_arg(s: str | None) -> tuple[int, int]:
    if s:
        y, m = (int(x) for x in s.split("-"))
        return y, m
    today = date.today()
    nxt = date(today.year, today.month, 28) + timedelta(days=7)
    return nxt.year, nxt.month


def _line_item(geo_id, unit_id, sizes, y, m, video=False):
    last = calendar.monthrange(y, m)[1]
    li = {
        "name": "avails forecast probe",
        "startDateTime": {"date": {"year": y, "month": m, "day": 1},
                          "hour": 0, "minute": 0, "second": 0, "timeZoneId": TZ},
        "endDateTime": {"date": {"year": y, "month": m, "day": last},
                        "hour": 23, "minute": 59, "second": 0, "timeZoneId": TZ},
        "lineItemType": "STANDARD",
        "costType": "CPM",
        "primaryGoal": {"goalType": "LIFETIME", "unitType": "IMPRESSIONS", "units": 1},
        "creativePlaceholders": [
            {"size": {"width": w, "height": h, "isAspectRatio": False}} for w, h in sizes
        ],
        "targeting": {
            "inventoryTargeting": {
                "targetedAdUnits": [{"adUnitId": str(unit_id), "includeDescendants": True}]},
            "geoTargeting": {"targetedLocations": [{"id": str(geo_id)}]},
        },
    }
    if video:
        # All three are required together — see module docstring.
        li["environmentType"] = "VIDEO_PLAYER"
        li["targeting"]["requestPlatformTargeting"] = {"targetedRequestPlatforms": ["VIDEO_PLAYER"]}
        li["videoMaxDuration"] = VIDEO_MAX_DURATION_MS
    return li


def _forecast(fc, li, attempts: int = 6):
    """One availability forecast, with backoff on GAM's forecast throttle."""
    for i in range(attempts):
        try:
            with _throttle:
                r = fc.getAvailabilityForecast(
                    {"lineItem": li},
                    {"includeTargetingCriteriaBreakdown": False,
                     "includeContendingLineItems": False},
                )
            return int(r.matchedUnits), int(r.availableUnits)
        except Exception as ex:  # noqa: BLE001
            msg = str(ex)
            # GAM names its forecast throttle EXCEEDED_QUOTA; the earlier
            # "QuotaExceeded" spelling never matched, so quota faults fell
            # straight through as hard failures.
            transient = any(k in msg for k in
                            ("EXCEEDED_QUOTA", "QuotaExceeded", "RATE_EXCEEDED",
                             "SERVER_ERROR", "Timeout", "timed out", "INTERNAL"))
            if not transient or i == attempts - 1:
                raise
            time.sleep(min(2 ** i * 5, 90))
    raise RuntimeError("unreachable")


def _resolve_units(soap) -> dict[str, int]:
    from googleads import ad_manager
    net = soap.GetService("NetworkService", version=API_VERSION).getCurrentNetwork()
    inv = soap.GetService("InventoryService", version=API_VERSION)
    stmt = (ad_manager.StatementBuilder(version=API_VERSION)
            .Where("parentId = :p")
            .WithBindVariable("p", int(net["effectiveRootAdUnitId"]))
            .Limit(200))
    units = {u.name: int(u.id) for u in inv.getAdUnitsByStatement(stmt.ToStatement()).results}
    missing = {DISPLAY_UNIT, VIDEO_UNIT} - units.keys()
    if missing:
        raise SystemExit(f"top-level ad unit(s) not found: {sorted(missing)}")
    return units


def _resolve_geos(soap) -> dict[str, tuple[int, str]]:
    """ISO country code -> (geo target id, canonical name)."""
    from googleads import ad_manager
    pql = soap.GetService("PublisherQueryLanguageService", version=API_VERSION)
    out, offset = {}, 0
    while True:
        q = (ad_manager.StatementBuilder(version=API_VERSION)
             .Select("Id, Name, CountryCode").From("Geo_Target")
             .Where("Type = 'Country' AND Targetable = true")
             .Limit(500).Offset(offset))
        res = pql.select(q.ToStatement())
        rows = getattr(res, "rows", None) or []
        for r in rows:
            gid, name, code = (v.value for v in r.values)
            if code:
                out[str(code)] = (int(gid), str(name))
        if len(rows) < 500:
            break
        offset += 500
    return out


def _recent_countries(client: GAMClient) -> pd.DataFrame:
    """Countries with real traffic in the trailing 28 days, by ISO code."""
    end = date.today() - timedelta(days=1)
    df = client._run_report(
        dimensions=["COUNTRY_CODE", "COUNTRY_NAME"],
        metrics=["IMPRESSIONS", "UNFILLED_IMPRESSIONS"],
        start_date=end - timedelta(days=27), end_date=end,
        filters=[("AD_UNIT_NAME_TOP_LEVEL", "IN", [DISPLAY_UNIT, VIDEO_UNIT])],
    )
    df["recent_avails"] = df.impressions + df.unfilled_impressions
    df = df[(df.country_code != "") & (~df.country_code.isin(EXCLUDE_CODES))]
    return (df.groupby(["country_code", "country_name"], as_index=False)["recent_avails"].sum()
              .sort_values("recent_avails", ascending=False))


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--month", help="month to forecast, YYYY-MM (default: next month)")
    ap.add_argument("--min-recent", type=int, default=1000,
                    help="skip countries under this many trailing-28d avails (default 1000)")
    ap.add_argument("--no-per-size", action="store_true",
                    help="only the de-duplicated display figure, not the four per-size calls")
    ap.add_argument("--workers", type=int, default=3)
    ap.add_argument("--resume-from", help="CSV from a prior run; already-computed cuts are skipped")
    ap.add_argument("--csv-dir")
    ap.add_argument("--xlsx")
    ap.add_argument("--top", type=int, default=25)
    args = ap.parse_args()

    y, m = _month_arg(args.month)
    label = f"{y:04d}-{m:02d}"
    if date(y, m, calendar.monthrange(y, m)[1]) < date.today():
        print(f"!! {label} is in the past — forecasts only run forward. "
              f"Use scripts/gam_intl_avails.py for historical actuals.")
        return 2
    print(f"# GAM non-US AVAILS FORECAST — {label} (ForecastService {API_VERSION})")
    print(f"# Display: `{DISPLAY_UNIT}` @ {', '.join(f'{w}x{h}' for w, h in DISPLAY_SIZES)}")
    print(f"# Video:   `{VIDEO_UNIT}` @ {VIDEO_SIZE[0]}x{VIDEO_SIZE[1]}\n")

    client = GAMClient()
    soap = client._get_soap_client()
    fc = soap.GetService("ForecastService", version=API_VERSION)

    units = _resolve_units(soap)
    geos = _resolve_geos(soap)
    countries = _recent_countries(client)
    countries = countries[countries.recent_avails >= args.min_recent]
    countries = countries[countries.country_code.isin(geos.keys())]
    print(f"{len(countries)} non-US countries with >={args.min_recent:,} trailing-28d avails "
          f"and a targetable Geo_Target\n")

    jobs = []
    for row in countries.itertuples():
        gid, _ = geos[row.country_code]
        jobs.append((row.country_name, "Display", ANY_SIZE, gid, units[DISPLAY_UNIT], DISPLAY_SIZES, False))
        jobs.append((row.country_name, "Video", f"{VIDEO_SIZE[0]}x{VIDEO_SIZE[1]}",
                     gid, units[VIDEO_UNIT], [VIDEO_SIZE], True))
        if not args.no_per_size:
            for w, h in DISPLAY_SIZES:
                jobs.append((row.country_name, "Display", f"{w}x{h}", gid,
                             units[DISPLAY_UNIT], [(w, h)], False))

    if args.resume_from and Path(args.resume_from).exists():
        prior = pd.read_csv(args.resume_from)
        have = {(r.country, r["format"], r["size"]) for _, r in prior.iterrows()}
        before = len(jobs)
        jobs = [j for j in jobs if (j[0], j[1], j[2]) not in have]
        print(f"resume: {before - len(jobs)} cuts already present in {args.resume_from}, "
              f"{len(jobs)} left to run")
    else:
        prior = None

    print(f"running {len(jobs)} forecast calls with {args.workers} workers …")
    rows, errors, done = [], [], 0
    t0 = time.time()
    with ThreadPoolExecutor(max_workers=args.workers) as ex:
        futs = {
            ex.submit(_forecast, fc, _line_item(gid, uid, sizes, y, m, video)): (cty, fmt, size)
            for cty, fmt, size, gid, uid, sizes, video in jobs
        }
        for fut in as_completed(futs):
            cty, fmt, size = futs[fut]
            done += 1
            if done % 100 == 0:
                print(f"  {done}/{len(jobs)} ({time.time()-t0:.0f}s)")
            try:
                matched, available = fut.result()
                rows.append({"month": label, "country": cty, "format": fmt, "size": size,
                             "matched": matched, "available": available})
            except Exception as exc:  # noqa: BLE001
                errors.append((cty, fmt, size, str(exc)[:120]))

    print(f"done in {time.time()-t0:.0f}s — {len(rows)} ok, {len(errors)} failed")
    for e in errors[:10]:
        print("  FAILED:", e)
    df = pd.DataFrame(rows)
    if prior is not None:
        df = pd.concat([prior, df], ignore_index=True) if not df.empty else prior
    if df.empty:
        raise SystemExit("no forecasts returned")
    # No fill_value: a cut that failed must stay NaN rather than read as
    # zero inventory, which is a different and very wrong claim.
    wide = df.pivot_table(index="country", columns=["format", "size"], values="available",
                          aggfunc="sum")
    order = ([("Display", f"{w}x{h}") for w, h in DISPLAY_SIZES]
             + [("Display", ANY_SIZE), ("Video", f"{VIDEO_SIZE[0]}x{VIDEO_SIZE[1]}")])
    wide = wide.reindex(columns=[c for c in order if c in wide.columns])
    sort_key = ("Display", ANY_SIZE)
    if sort_key in wide.columns:
        wide = wide.sort_values(sort_key, ascending=False)

    print(f"\n## {label} non-US totals (available = unreserved, sellable)\n")
    print(df.groupby(["format", "size"])[["matched", "available"]].sum().to_string())
    print(f"\n## {label} available avails by country — top {args.top}\n")
    print(wide.head(args.top).to_string())
    print(f"\n… {max(0, len(wide) - args.top)} more countries ({len(wide)} total)")
    gaps = int(wide.isna().sum().sum())
    if gaps:
        print(f"\n!! {gaps} cut(s) across {int(wide.isna().any(axis=1).sum())} countries did NOT "
              f"return a forecast (blank, NOT zero). Re-run with "
              f"--resume-from <csv> to fill them.")
    print("\nNOTE: display sizes OVERLAP — one opportunity is eligible for several sizes, "
          "so the four size columns must not be summed. Use the de-duplicated column.")

    if args.csv_dir:
        out = Path(args.csv_dir); out.mkdir(parents=True, exist_ok=True)
        df.to_csv(out / f"gam_avails_forecast_{label}.csv", index=False)
        print(f"\nwrote {out}/gam_avails_forecast_{label}.csv")
    if args.xlsx:
        p = Path(args.xlsx); p.parent.mkdir(parents=True, exist_ok=True)
        flat = wide.copy()
        flat.columns = [f"{f} {s}" for f, s in flat.columns]
        with pd.ExcelWriter(p, engine="openpyxl") as xl:
            pd.DataFrame({f"Newsweek — non-US avails FORECAST, {label}": [
                f"Source: GAM ForecastService {API_VERSION} (getAvailabilityForecast), "
                f"network {os.environ.get('GAM_NETWORK_ID','')}. Run {date.today()}.",
                "This is a FORECAST of a month that has not happened, not a report. It carries "
                "GAM's own traffic model and nets out inventory already reserved by other line items.",
                f"Display = top-level ad unit `{DISPLAY_UNIT}` limited to "
                f"{', '.join(f'{w}x{h}' for w, h in DISPLAY_SIZES)}. Video = `{VIDEO_UNIT}`.",
                "Available = unreserved inventory you can still sell. Matched = the gross pool "
                "before other line items' reservations. Quote Available.",
                "DISPLAY SIZES OVERLAP AND MUST NOT BE SUMMED — one opportunity is eligible for "
                "several sizes at once. Use the de-duplicated column for a single figure.",
                "United States is excluded.",
            ]}).to_excel(xl, sheet_name="Read me", index=False)
            df.groupby(["format", "size"], as_index=False)[["matched", "available"]].sum() \
              .to_excel(xl, sheet_name="Summary", index=False)
            flat.reset_index().rename(columns={"country": "Country"}) \
              .to_excel(xl, sheet_name=f"By country {label}", index=False)
            from openpyxl.styles import Alignment, Font
            from openpyxl.utils import get_column_letter
            for name in xl.book.sheetnames:
                ws = xl.book[name]
                ws.freeze_panes = "A2"
                for cell in ws[1]:
                    cell.font = Font(bold=True)
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
                for col in ws.columns:
                    width = max((len(str(c.value)) for c in col if c.value is not None), default=10)
                    ws.column_dimensions[get_column_letter(col[0].column)].width = \
                        min(max(width + 2, 12), 95 if name == "Read me" else 26)
                if name != "Read me":
                    for r in ws.iter_rows(min_row=2):
                        for cell in r:
                            if isinstance(cell.value, (int, float)):
                                cell.number_format = "#,##0"
        print(f"wrote {p}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
