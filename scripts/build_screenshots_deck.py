#!/usr/bin/env python3
"""Build the client-facing proof-of-placement deck as a PowerPoint file.

The deliverable is always a .pptx (Roger, 2026-09-23): cover -> campaign
summary -> one slide per IN-CONTEXT shot. No close-up crops, and one deck per
advertiser. See docs/screenshots_document.md.

Input is a small JSON spec, filled from the pull (`pull_screenshots_source.yml`)
and the capture artifacts (`capture_screenshots.yml`):

    {
      "advertiser": "Elevance Health",
      "campaign": "AI Health Summit 2026 · Display",
      "captured": "23 Sep 2026, 11:18 AM ET",
      "page_url": "newsweek.com/could-diet-reduce-alzheimers-risk-...",
      "summary_note": "The campaign runs across newsweek.com; ...",
      "facts": [["Advertiser", "Elevance Health"], ["Flight", "..."], ...],
      "shots": [
        {"size": "970x250", "device": "Desktop", "title": "Billboard in the article body",
         "file": "shots/7440274079_138612841039_970x250_desktop_context.png",
         "alt": "Elevance Health 970x250 billboard in a Newsweek Health article"}
      ]
    }

`file` paths resolve relative to the spec. Screenshots are scaled to fit their
panel, never cropped: the proof is the whole point.

Usage:
  python scripts/build_screenshots_deck.py spec.json "Elevance Health - Proof of Placement.pptx"
"""
from __future__ import annotations

import argparse
import json
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Emu, Inches, Pt

# Newsweek "Paper" (same values as the dashboard tokens): warm paper, ink, and
# brand red as chrome only — the eyebrow tick, never data.
PAPER, PANEL, INK, SOFT, RED = "FEFCF6", "F1ECE0", "1F1E19", "55524A", "E91D0C"
ON_INK_MUTED, ON_INK_FOOT, ON_INK_SUB, RULE = "C9C3B3", "A9A393", "D8D2C2", "DDD6C6"
SERIF, SANS = "Cambria", "Calibri"  # ship with Office, so the file renders as built
W, H, M = 13.333, 7.5, 0.6


def _rgb(h: str) -> RGBColor:
    return RGBColor.from_string(h)


def _text(slide, text, x, y, w, h, *, size, color, font=SANS, bold=False,
          spacing=None):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    run = tf.paragraphs[0].add_run()
    run.text = text
    f = run.font
    f.name, f.size, f.bold, f.color.rgb = font, Pt(size), bold, _rgb(color)
    if spacing is not None:  # character spacing, in hundredths of a point
        run._r.get_or_add_rPr().set("spc", str(int(spacing * 100)))
    return tb


def _rect(slide, x, y, w, h, color, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = _rgb(color)
    s.line.fill.background()
    s.shadow.inherit = False
    return s


def _background(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = _rgb(color)


def _eyebrow(slide, text, on_ink=False):
    _rect(slide, M, 0.73, 0.35, 0.04, RED)
    _text(slide, text.upper(), M + 0.5, 0.6, 10, 0.3, size=12, bold=True,
          spacing=3, color=ON_INK_MUTED if on_ink else SOFT)


def _footer(slide, text, on_ink=False):
    _text(slide, text, M, 6.95, W - 2 * M, 0.3, size=11,
          color=ON_INK_FOOT if on_ink else SOFT)


def build(spec: dict, base: Path, out: Path) -> None:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    blank = prs.slide_layouts[6]
    prs.core_properties.title = f"{spec['advertiser']} — Proof of Placement"
    prs.core_properties.author = "Newsweek"

    # Cover
    s = prs.slides.add_slide(blank)
    _background(s, INK)
    _eyebrow(s, "Newsweek · Proof of placement", on_ink=True)
    _text(s, spec["advertiser"], M, 2.5, W - 2 * M, 1.3, font=SERIF, size=60,
          bold=True, color=PAPER)
    _text(s, spec["campaign"], M, 3.85, W - 2 * M, 0.6, size=24, color=ON_INK_SUB)
    _footer(s, f"Captured {spec['captured']} · newsweek.com", on_ink=True)

    # Campaign summary
    s = prs.slides.add_slide(blank)
    _background(s, PAPER)
    _eyebrow(s, "Campaign")
    _text(s, "What ran, and where we shot it", M, 1.0, W - 2 * M, 0.7,
          font=SERIF, size=32, bold=True, color=INK)
    facts = spec["facts"]
    tbl = s.shapes.add_table(len(facts), 2, Inches(M), Inches(2.0),
                             Inches(W - 2 * M), Inches(0.52 * len(facts))).table
    tbl.first_row = False
    tbl.horz_banding = False
    tbl.columns[0].width = Inches(3.2)
    tbl.columns[1].width = Inches(W - 2 * M - 3.2)
    for r, (k, v) in enumerate(facts):
        tbl.rows[r].height = Inches(0.52)
        for c, (val, color, bold) in enumerate(((k, SOFT, True), (v, INK, False))):
            cell = tbl.cell(r, c)
            cell.fill.solid()
            cell.fill.fore_color.rgb = _rgb(PAPER if r % 2 else PANEL)
            cell.margin_left = cell.margin_right = Inches(0.14)
            tf = cell.text_frame
            tf.text = ""
            run = tf.paragraphs[0].add_run()
            run.text = val
            run.font.name, run.font.size = SANS, Pt(15)
            run.font.bold, run.font.color.rgb = bold, _rgb(color)
    _footer(s, spec["summary_note"])

    # One slide per in-context shot
    shots = spec["shots"]
    for i, sh in enumerate(shots, 1):
        s = prs.slides.add_slide(blank)
        _background(s, PAPER)
        _eyebrow(s, f"{sh['size']} · {sh['device']} · In context")
        _text(s, sh["title"], M, 1.0, W - 2 * M, 0.7, font=SERIF, size=30,
              bold=True, color=INK)
        px, py, pw, ph, pad = M, 1.9, W - 2 * M, 4.85, 0.15
        _rect(s, px, py, pw, ph, PANEL, MSO_SHAPE.ROUNDED_RECTANGLE).adjustments[0] = 0.02
        img = base / sh["file"]
        with Image.open(img) as im:
            iw, ih = im.size
        bw, bh = pw - 2 * pad, ph - 2 * pad
        k = min(bw / iw, bh / ih)  # contain: scale to fit, never crop
        w, h = iw * k, ih * k
        pic = s.shapes.add_picture(str(img), Inches(px + pad + (bw - w) / 2),
                                   Inches(py + pad + (bh - h) / 2),
                                   Inches(w), Inches(h))
        pic._element.nvPicPr.cNvPr.set("descr", sh.get("alt", ""))
        _footer(s, f"{spec['page_url']} · {spec['captured']} · {sh['size']} · "
                   f"{i}/{len(shots)}")

    prs.save(str(out))


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("spec", type=Path)
    ap.add_argument("out", type=Path)
    args = ap.parse_args()
    spec = json.loads(args.spec.read_text())
    build(spec, args.spec.parent, args.out)
    print(f"wrote {args.out}  ({2 + len(spec['shots'])} slides)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
